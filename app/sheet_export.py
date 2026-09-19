"""Luckysheet の JSON を .xlsx に書き戻す（ブラウザ編集の保存）。

値・数式・結合・列幅/行高・基本書式（太字/斜体/文字色/背景/フォントサイズ/配置/表示形式）を保持する。
数式のキャッシュ値は openpyxl では書けないため、LibreOffice があれば再計算して焼き込む。
"""
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def _color(v: Any) -> str | None:
    if not v or not isinstance(v, str) or not v.startswith("#"):
        return None
    return "FF" + v[1:7].upper() if len(v) >= 7 else None


def _cell_value(v: dict) -> Any:
    if v.get("f"):
        f = str(v["f"])
        return f if f.startswith("=") else "=" + f
    ct = v.get("ct") or {}
    # リッチテキスト（luckyexcel の取り込みでは文字列が inlineStr の runs に入る）
    if ct.get("t") == "inlineStr" and isinstance(ct.get("s"), list):
        text = "".join(str(run.get("v", "")) for run in ct["s"] if isinstance(run, dict))
        if text:
            return text
    val = v.get("v")
    if val is None or val == "":
        return v.get("m") or None
    if ct.get("t") == "n" or isinstance(val, (int, float)):
        try:
            num = float(val)
            return int(num) if num.is_integer() else num
        except (TypeError, ValueError):
            return val
    return val


def _apply_style(cell, v: dict) -> None:
    font_kw = {}
    if v.get("bl"):
        font_kw["bold"] = True
    if v.get("it"):
        font_kw["italic"] = True
    if v.get("fs"):
        try:
            font_kw["size"] = float(v["fs"])
        except (TypeError, ValueError):
            pass
    if c := _color(v.get("fc")):
        font_kw["color"] = c
    if font_kw:
        cell.font = Font(**font_kw)
    if bg := _color(v.get("bg")):
        cell.fill = PatternFill("solid", fgColor=bg)
    ht = {0: "center", 1: "left", 2: "right"}.get(v.get("ht"))
    vt = {0: "center", 1: "top", 2: "bottom"}.get(v.get("vt"))
    if ht or vt or v.get("tb") == "2":
        cell.alignment = Alignment(horizontal=ht, vertical=vt, wrap_text=(v.get("tb") == "2") or None)
    fa = (v.get("ct") or {}).get("fa")
    if fa and fa != "General":
        cell.number_format = fa


def luckysheet_to_workbook(sheets: list[dict]) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    for s in sorted(sheets, key=lambda x: x.get("order", 0)):
        ws = wb.create_sheet(title=str(s.get("name") or "Sheet1")[:31])
        cells = s.get("celldata")
        if not cells and s.get("data"):
            cells = [{"r": r, "c": c, "v": v} for r, row in enumerate(s["data"]) for c, v in enumerate(row) if v]
        for cd in cells or []:
            v = cd.get("v")
            if not isinstance(v, dict):
                if v is not None:
                    ws.cell(row=cd["r"] + 1, column=cd["c"] + 1, value=v)
                continue
            val = _cell_value(v)
            cell = ws.cell(row=cd["r"] + 1, column=cd["c"] + 1)
            if val is not None:
                cell.value = val
            _apply_style(cell, v)
        cfg = s.get("config") or {}
        for m in (cfg.get("merge") or {}).values():
            try:
                ws.merge_cells(start_row=m["r"] + 1, start_column=m["c"] + 1,
                               end_row=m["r"] + m["rs"], end_column=m["c"] + m["cs"])
            except Exception:
                pass
        for c, w in (cfg.get("columnlen") or {}).items():
            try:
                ws.column_dimensions[get_column_letter(int(c) + 1)].width = max(4, float(w) / 7)
            except Exception:
                pass
        for r, h in (cfg.get("rowlen") or {}).items():
            try:
                ws.row_dimensions[int(r) + 1].height = float(h) * 0.75
            except Exception:
                pass
    return wb


def bake(path: Path) -> bool:
    """LibreOffice で開き直して数式のキャッシュ値を焼き込む。soffice が無ければ False"""
    if not shutil.which("soffice"):
        return False
    with tempfile.TemporaryDirectory() as td:
        r = subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", td, str(path)],
                           capture_output=True, timeout=120)
        out = Path(td) / path.name
        if r.returncode != 0 or not out.exists():
            return False
        shutil.move(str(out), str(path))
    return True


def save_xlsx(sheets: list[dict], path: Path) -> bool:
    wb = luckysheet_to_workbook(sheets)
    wb.save(path)
    return bake(path)
