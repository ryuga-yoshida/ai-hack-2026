"""架空企業「青葉ビバレッジ株式会社」の売上見込表を2版生成する。

openpyxl が書いたファイルは数式のキャッシュ値を持たないため、
生成後に LibreOffice で再保存してキャッシュ値を焼き込む（下記 bake）。

    python fixtures/excel/make_fixtures.py
"""
import shutil
import subprocess
import tempfile
from pathlib import Path

from openpyxl import Workbook

HERE = Path(__file__).parent
HEADER = ["商品名", "第1四半期", "第2四半期", "第3四半期", "第4四半期"]

# v1: 9/10 時点。商品A の第3四半期は 1200
V1 = [
    ["商品B", 800, 820, 850, 900],
    ["商品C", 500, 500, 500, 500],
    ["商品E", 300, 320, 340, 360],
    ["商品A", 1000, 1100, 1200, 1300],
    ["商品F", 200, 210, 220, 230],
    ["商品G", 150, 150, 160, 160],
    ["商品H", 400, 420, 440, 460],
    ["商品I", 120, 120, 130, 130],
    ["商品J", 90, 95, 100, 105],
    ["商品K", 60, 60, 65, 65],
]

# v2: 9/18 山田が更新。
#   - 商品D を先頭に挿入（9/8 の決定通り。以降の行がズレる）
#   - 商品C を削除
#   - 商品B の第1四半期 800 → 850（9/8 の決定通り。検知してはいけない）
#   - 商品A の第3四半期 1200 → 1500（9/15 の「据え置き」決定と矛盾。主役）
#   - 合計の数式範囲を修正（末尾行が漏れていた）
V2 = [
    ["商品D", 0, 0, 300, 400],
    ["商品B", 850, 820, 850, 900],
    ["商品E", 300, 320, 340, 360],
    ["商品A", 1000, 1100, 1500, 1300],
    ["商品F", 200, 210, 220, 230],
    ["商品G", 150, 150, 160, 160],
    ["商品H", 400, 420, 440, 460],
    ["商品I", 120, 120, 130, 130],
    ["商品J", 90, 95, 100, 105],
    ["商品K", 60, 60, 65, 65],
]


def build(rows, sum_last_row, path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "売上見込"
    ws.append(HEADER)
    for r in rows:
        ws.append(r)
    total_row = len(rows) + 2
    ws.cell(row=total_row, column=1, value="合計")
    for col in "BCDE":
        ws[f"{col}{total_row}"] = f"=SUM({col}2:{col}{sum_last_row})"
    wb.save(path)


def bake(path: Path):
    """LibreOffice で開き直して数式のキャッシュ値を焼き込む"""
    with tempfile.TemporaryDirectory() as td:
        subprocess.run(
            ["soffice", "--headless", "--convert-to", "xlsx", "--outdir", td, str(path)],
            check=True, capture_output=True,
        )
        shutil.move(Path(td) / path.name, path)


if __name__ == "__main__":
    v1, v2 = HERE / "売上見込_v1.xlsx", HERE / "売上見込_v2.xlsx"
    build(V1, sum_last_row=10, path=v1)   # 末尾（商品K, 11行目）が漏れている
    build(V2, sum_last_row=11, path=v2)   # 修正済み
    for p in (v1, v2):
        bake(p)
        print("wrote", p)
